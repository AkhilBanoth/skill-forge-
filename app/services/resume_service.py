
from io import BytesIO

from docx import Document
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.models.skill import Skill


MAX_RESUME_SIZE = 5 * 1024 * 1024

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


async def extract_resume_text(file: UploadFile) -> str:
    filename = (file.filename or "").lower()
    extension = "." + filename.rsplit(".", 1)[-1] if "." in filename else ""

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF and DOCX resume files are supported.",
        )

    content = await file.read(MAX_RESUME_SIZE + 1)

    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty.",
        )

    if len(content) > MAX_RESUME_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Resume file must be 5 MB or smaller.",
        )

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Password-protected PDFs are not supported.",
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        else:
            document = Document(BytesIO(content))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            table_text = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]

            text = "\n".join(paragraphs + table_text)

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not read this file. Please upload a valid PDF or DOCX.",
        )

    text = text.strip()

    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No readable text was found. "
                "Scanned PDFs may require OCR."
            ),
        )

    return text


def detect_resume_skills(
    db: Session,
    resume_text: str,
) -> list[dict]:
    known_skills = db.query(Skill).all()
    normalized_text = resume_text.casefold()
    detected = []

    for skill in known_skills:
        skill_name = skill.name.strip()

        if not skill_name:
            continue

        import re

        pattern = (
            r"(?<![a-z0-9+#.])"
            + re.escape(skill_name.casefold())
            + r"(?![a-z0-9+#.])"
        )

        if re.search(pattern, normalized_text):
            detected.append(
                {
                    "skill_id": skill.id,
                    "skill_name": skill.name,
                    "category": skill.category,
                }
            )

    return detected
